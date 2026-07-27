"""Extract content brief docx to UTF-8 text."""
from pathlib import Path

try:
    from docx import Document
except ImportError:
    import subprocess
    subprocess.check_call(['pip', 'install', 'python-docx', '-q'])
    from docx import Document

doc = Document(r'c:\Users\Dev\Downloads\klugminds-website-content-brief_v1.docx')
out = Path(__file__).resolve().parents[1] / 'tmp-content-brief.txt'

lines: list[str] = []
for p in doc.paragraphs:
    if p.text.strip():
        lines.append(p.text)

for table in doc.tables:
    lines.append('--- TABLE ---')
    for row in table.rows:
        lines.append(' | '.join(cell.text.strip() for cell in row.cells))

out.write_text('\n'.join(lines), encoding='utf-8')
print(f'Wrote {len(lines)} lines to {out}')
